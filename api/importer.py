"""
importer.py — Parse Excel EHR → insère dans Turso via libSQL.
"""
import re, hashlib
from datetime import datetime
import libsql_client
from db import get_client, rows_to_dicts

try:
    import xlrd
    HAVE_XLRD = True
except ImportError:
    HAVE_XLRD = False

try:
    import openpyxl
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

TEAM_ROW = 11
COACH_ROW = 13
DATE_COL  = 4
FIRST_TEAM_COL = 5
HEADER_ROWS = 14

def _safe(v) -> str:
    return str(v).strip() if v is not None else ""

def _parse_date_iso(s: str) -> str:
    try:
        p = s.split(" ")[0].split("/")
        d, m, y = p[0], p[1], p[2]
        if len(y) == 2: y = "20" + y
        return f"{y}-{m.zfill(2)}-{d.zfill(2)}"
    except Exception:
        return "9999-99-99"

def _extract_time(t: str) -> str:
    m = re.search(r"à\s*([\d]+h[\d]*)", t, re.IGNORECASE)
    return m.group(1) if m else ""

def _extract_opponent(text: str) -> str:
    line = text.split("\n")[0].strip()
    if " - EHR" in line: return line.split(" - EHR")[0].strip()
    if "EHR - " in line: return line.split("EHR - ")[1].strip()
    return line.replace("EHR", "").replace("-", "").strip()

def _is_home(text: str):
    line = text.split("\n")[0]
    if re.search(r"^EHR[\s\-]", line): return 1
    if re.search(r"\-\s*EHR", line): return 0
    return None

def _match_type(text: str, journee: str) -> str:
    t, j = text.lower(), journee.lower()
    if re.search(r"exempt|forfait", t): return "exempt"
    if re.search(r"report", t) and not re.search(r"coupe|cdf", t+j): return "report"
    if re.search(r"coupe|cdf|moselle|moelle", t+j): return "coupe"
    if re.search(r"amical|tournoi|plateau", t+j): return "amical"
    return "champ"

def _import_key(date_str, team, text) -> str:
    raw = f"{date_str}|{team}|{text.split(chr(10))[0]}"
    return hashlib.md5(raw.encode()).hexdigest()[:16]

def _is_noise(text: str) -> bool:
    noise = [r"^journée\s+\d", r"^j\d+$", r"^coupe de", r"^match amical$",
             r"^forfait g", r"^retrait ", r"impératif", r"report du",
             r"avancé au", r"modif salle", r"^\d+h\d*\s*[-–]", r"^à placer"]
    t = text.lower().strip()
    return any(re.search(p, t) for p in noise) or len(t) < 4

def _read_xls(path):
    wb = xlrd.open_workbook(path)
    sh = wb.sheet_by_name("Planning")
    teams = {}
    for c in range(FIRST_TEAM_COL, sh.ncols):
        name = _safe(sh.cell_value(TEAM_ROW, c))
        coach = _safe(sh.cell_value(COACH_ROW, c))
        if name: teams[c] = {"name": name, "coach": coach}
    rows = []
    for r in range(HEADER_ROWS, sh.nrows):
        rows.append([_safe(sh.cell_value(r, c)) for c in range(sh.ncols)])
    return teams, rows

def _read_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = wb["Planning"]
    teams = {}
    for c in range(FIRST_TEAM_COL, sh.max_column + 1):
        name = _safe(sh.cell(row=TEAM_ROW+1, column=c+1).value)
        coach = _safe(sh.cell(row=COACH_ROW+1, column=c+1).value)
        if name: teams[c] = {"name": name, "coach": coach}
    rows = []
    for r in range(HEADER_ROWS+1, sh.max_row+1):
        rows.append([_safe(sh.cell(row=r, column=c+1).value) for c in range(sh.max_column)])
    return teams, rows

def parse_excel(path: str) -> list:
    if path.endswith(".xls") and HAVE_XLRD:
        teams, rows = _read_xls(path)
    elif HAVE_OPENPYXL:
        teams, rows = _read_xlsx(path)
    else:
        raise RuntimeError("xlrd ou openpyxl requis")

    matches, current_date = [], ""
    for row in rows:
        date_val = row[DATE_COL] if len(row) > DATE_COL else ""
        if date_val and re.search(r"\d{2}/\d{2}/\d{2}", date_val):
            current_date = date_val.replace("\n", " ").strip()
        if not current_date: continue
        for col, team in teams.items():
            if col >= len(row): continue
            val = row[col]
            if not val or val in ("0.0", " ") or _is_noise(val): continue
            matches.append({
                "date_str": current_date,
                "date_iso": _parse_date_iso(current_date),
                "team_name": team["name"],
                "coach": team["coach"],
                "match_text": val,
                "opponent": _extract_opponent(val),
                "home": _is_home(val),
                "time_str": _extract_time(val),
                "match_type": _match_type(val, ""),
                "excel_import_id": _import_key(current_date, team["name"], val),
            })
    return matches

def import_to_turso(path: str, filename: str) -> dict:
    now = datetime.now().isoformat()
    try:
        parsed = parse_excel(path)
    except Exception as e:
        return {"status": "error", "message": str(e), "created": 0, "updated": 0, "skipped": 0}

    created = updated = skipped = 0

    with get_client() as client:
        for m in parsed:
            key = m["excel_import_id"]
            existing = client.execute(
                libsql_client.Statement("SELECT id, manually_edited FROM matches WHERE excel_import_id = ?", [key])
            )
            rows = list(existing.rows)
            if rows:
                match_id, manually_edited = rows[0][0], rows[0][1]
                if manually_edited:
                    skipped += 1
                    continue
                client.execute(libsql_client.Statement(
                    "UPDATE matches SET match_text=?, opponent=?, home=?, time_str=?, match_type=?, updated_at=? WHERE id=?",
                    [m["match_text"], m["opponent"], m["home"], m["time_str"], m["match_type"], now, match_id]
                ))
                updated += 1
            else:
                client.execute(libsql_client.Statement(
                    """INSERT INTO matches
                       (date_str,date_iso,team_name,coach,match_text,opponent,home,time_str,match_type,
                        note,manually_edited,excel_import_id,created_at,updated_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    [m["date_str"], m["date_iso"], m["team_name"], m["coach"], m["match_text"],
                     m["opponent"], m["home"], m["time_str"], m["match_type"],
                     "", 0, key, now, now]
                ))
                created += 1

        client.execute(libsql_client.Statement(
            "INSERT INTO import_logs (filename,imported_at,rows_created,rows_updated,rows_skipped,status,message) VALUES (?,?,?,?,?,?,?)",
            [filename, now, created, updated, skipped, "ok", ""]
        ))

    return {"status": "ok", "filename": filename, "created": created, "updated": updated, "skipped": skipped, "message": ""}
